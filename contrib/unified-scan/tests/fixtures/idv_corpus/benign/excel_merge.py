import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook


def merge(folder: str, target: str) -> None:
    out = Workbook()
    ws_out = out.active
    for f in sorted(Path(folder).glob("*.xlsx")):
        wb = load_workbook(f, read_only=True)
        for row in wb.active.iter_rows(values_only=True):
            ws_out.append(row)
    out.save(target)


if __name__ == "__main__":
    merge(sys.argv[1], sys.argv[2])
